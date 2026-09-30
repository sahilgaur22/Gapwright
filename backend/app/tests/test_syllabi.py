import io
import uuid
from collections.abc import AsyncGenerator

import docx
import pytest
from httpx import ASGITransport, AsyncClient
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.core.security import create_access_token
from app.db.base import Base
from app.db.models.user import Institution, User
from app.db.session import get_db
from app.main import app


@pytest.fixture
async def test_session() -> AsyncGenerator[AsyncSession, None]:
    test_engine = create_async_engine("sqlite+aiosqlite:///:memory:", echo=False)
    async with test_engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    session_factory = async_sessionmaker(
        bind=test_engine, class_=AsyncSession, expire_on_commit=False
    )

    async with session_factory() as session:
        yield session

    async with test_engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)

    await test_engine.dispose()


@pytest.fixture
async def client(test_session: AsyncSession) -> AsyncGenerator[AsyncClient, None]:
    async def override_get_db() -> AsyncGenerator[AsyncSession, None]:
        yield test_session

    app.dependency_overrides[get_db] = override_get_db
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as c:
        yield c
    app.dependency_overrides.clear()


def _make_pdf_bytes(title: str, body: str) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.drawString(100, 750, title)
    c.drawString(100, 720, body)
    c.save()
    return buf.getvalue()


def _make_docx_bytes(title: str, topics: list[str]) -> bytes:
    doc = docx.Document()
    doc.add_paragraph(title)
    table = doc.add_table(rows=len(topics), cols=2)
    for i, topic in enumerate(topics):
        table.cell(i, 0).text = f"Unit {i + 1}"
        table.cell(i, 1).text = topic
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def _create_user(
    session: AsyncSession,
    email: str,
    *,
    role: str = "educator",
    institution_id: uuid.UUID | None = None,
) -> tuple[User, str]:
    user = User(
        email=email,
        password_hash="somehash",
        role=role,
        institution_id=institution_id,
    )
    session.add(user)
    await session.commit()
    await session.refresh(user)
    token = create_access_token(user.id)
    return user, token


@pytest.mark.asyncio
async def test_upload_syllabus_pdf(
    client: AsyncClient, test_session: AsyncSession
) -> None:
    inst = Institution(name="MIT", city="Cambridge")
    test_session.add(inst)
    await test_session.commit()
    await test_session.refresh(inst)

    _, token = await _create_user(test_session, "prof@mit.edu", institution_id=inst.id)
    pdf_bytes = _make_pdf_bytes(
        "CS 101: Introduction to Computer Science",
        "Covers Python, algorithms, and data structures."
    )

    response = await client.post(
        "/api/v1/syllabi",
        headers={"Authorization": f"Bearer {token}"},
        data={"title": "Intro to CS", "department": "Computer Science"},
        files={"file": ("cs101.pdf", pdf_bytes, "application/pdf")},
    )

    assert response.status_code == 201
    data = response.json()
    assert data["title"] == "Intro to CS"
    assert data["department"] == "Computer Science"
    assert data["filename"] == "cs101.pdf"
    assert data["status"] == "uploaded"
    assert "CS 101: Introduction to Computer Science" in data["raw_text"]
    assert "algorithms, and data structures" in data["raw_text"]
    assert data["institution_id"] == str(inst.id)


@pytest.mark.asyncio
async def test_upload_syllabus_docx(
    client: AsyncClient, test_session: AsyncSession
) -> None:
    _, token = await _create_user(test_session, "instructor@uni.edu")
    docx_bytes = _make_docx_bytes(
        "Cloud Computing Architecture",
        ["Docker Containers", "Kubernetes Orchestration", "Terraform IAC"],
    )

    response = await client.post(
        "/api/v1/syllabi",
        headers={"Authorization": f"Bearer {token}"},
        data={"title": "Cloud Computing", "department": "Information Systems"},
        files={
            "file": (
                "cloud.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 201
    data = response.json()
    assert data["title"] == "Cloud Computing"
    assert data["department"] == "Information Systems"
    assert "Cloud Computing Architecture" in data["raw_text"]
    assert "Kubernetes Orchestration" in data["raw_text"]


@pytest.mark.asyncio
async def test_upload_empty_file_rejected(
    client: AsyncClient, test_session: AsyncSession
) -> None:
    _, token = await _create_user(test_session, "user@uni.edu")
    response = await client.post(
        "/api/v1/syllabi",
        headers={"Authorization": f"Bearer {token}"},
        data={"title": "Empty File Test"},
        files={"file": ("empty.pdf", b"", "application/pdf")},
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()


@pytest.mark.asyncio
async def test_upload_unsupported_format_rejected(
    client: AsyncClient, test_session: AsyncSession
) -> None:
    _, token = await _create_user(test_session, "user2@uni.edu")
    response = await client.post(
        "/api/v1/syllabi",
        headers={"Authorization": f"Bearer {token}"},
        data={"title": "Text File"},
        files={"file": ("notes.txt", b"Plain text syllabus", "text/plain")},
    )
    assert response.status_code == 400
    assert "unsupported" in response.json()["detail"].lower()


@pytest.mark.asyncio
async def test_list_syllabi_scoped_to_institution(
    client: AsyncClient, test_session: AsyncSession
) -> None:
    # Create two institutions
    inst1 = Institution(name="Stanford University", city="Stanford")
    inst2 = Institution(name="UC Berkeley", city="Berkeley")
    test_session.add_all([inst1, inst2])
    await test_session.commit()
    await test_session.refresh(inst1)
    await test_session.refresh(inst2)

    _, token1 = await _create_user(
        test_session, "user1@stanford.edu", institution_id=inst1.id
    )
    _, token2 = await _create_user(
        test_session, "user2@berkeley.edu", institution_id=inst2.id
    )

    pdf1 = _make_pdf_bytes("Stanford CS 106A", "Programming Methodologies")
    pdf2 = _make_pdf_bytes("Berkeley CS 61A", "Structure and Interpretation")

    # Upload for inst1
    await client.post(
        "/api/v1/syllabi",
        headers={"Authorization": f"Bearer {token1}"},
        data={"title": "CS 106A", "department": "CS"},
        files={"file": ("cs106a.pdf", pdf1, "application/pdf")},
    )

    # Upload for inst2
    await client.post(
        "/api/v1/syllabi",
        headers={"Authorization": f"Bearer {token2}"},
        data={"title": "CS 61A", "department": "EECS"},
        files={"file": ("cs61a.pdf", pdf2, "application/pdf")},
    )

    # User 1 list
    res1 = await client.get(
        "/api/v1/syllabi", headers={"Authorization": f"Bearer {token1}"}
    )
    assert res1.status_code == 200
    data1 = res1.json()
    assert len(data1) == 1
    assert data1[0]["title"] == "CS 106A"

    # User 2 list
    res2 = await client.get(
        "/api/v1/syllabi", headers={"Authorization": f"Bearer {token2}"}
    )
    assert res2.status_code == 200
    data2 = res2.json()
    assert len(data2) == 1
    assert data2[0]["title"] == "CS 61A"


@pytest.mark.asyncio
async def test_get_syllabus_detail_and_cross_institution_isolation(
    client: AsyncClient, test_session: AsyncSession
) -> None:
    inst1 = Institution(name="Oxford", city="Oxford")
    inst2 = Institution(name="Cambridge", city="Cambridge")
    test_session.add_all([inst1, inst2])
    await test_session.commit()
    await test_session.refresh(inst1)
    await test_session.refresh(inst2)

    _, token1 = await _create_user(
        test_session, "user1@oxford.edu", institution_id=inst1.id
    )
    _, token2 = await _create_user(
        test_session, "user2@cambridge.edu", institution_id=inst2.id
    )

    pdf = _make_pdf_bytes("Oxford AI", "Logic and knowledge representation")
    upload_res = await client.post(
        "/api/v1/syllabi",
        headers={"Authorization": f"Bearer {token1}"},
        data={"title": "Oxford AI"},
        files={"file": ("ai.pdf", pdf, "application/pdf")},
    )
    syllabus_id = upload_res.json()["id"]

    # Inst1 user gets full details
    get_res1 = await client.get(
        f"/api/v1/syllabi/{syllabus_id}",
        headers={"Authorization": f"Bearer {token1}"},
    )
    assert get_res1.status_code == 200
    detail = get_res1.json()
    assert detail["id"] == syllabus_id
    assert "Logic and knowledge representation" in detail["raw_text"]

    # Inst2 user cannot access Inst1 syllabus
    get_res2 = await client.get(
        f"/api/v1/syllabi/{syllabus_id}",
        headers={"Authorization": f"Bearer {token2}"},
    )
    assert get_res2.status_code == 404


@pytest.mark.asyncio
async def test_endpoints_require_authentication(client: AsyncClient) -> None:
    res_list = await client.get("/api/v1/syllabi")
    assert res_list.status_code == 401

    res_get = await client.get(f"/api/v1/syllabi/{uuid.uuid4()}")
    assert res_get.status_code == 401

    res_post = await client.post("/api/v1/syllabi", data={"title": "Unauthorized"})
    assert res_post.status_code == 401
