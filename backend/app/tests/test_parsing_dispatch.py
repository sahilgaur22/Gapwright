import io

import docx
import pytest
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from app.services.parsing.dispatcher import detect_format, parse_document
from app.services.parsing.exceptions import UnsupportedFormatError


def _make_pdf_bytes(text: str) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.drawString(100, 750, text)
    c.save()
    return buf.getvalue()


def _make_docx_bytes(text: str) -> bytes:
    doc = docx.Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_detect_format_by_magic_bytes() -> None:
    pdf_bytes = _make_pdf_bytes("PDF Syllabus")
    docx_bytes = _make_docx_bytes("DOCX Syllabus")

    assert detect_format(pdf_bytes) == "pdf"
    assert detect_format(docx_bytes) == "docx"


def test_detect_format_by_content_type() -> None:
    dummy_bytes = b"non-magic-content"
    assert detect_format(dummy_bytes, content_type="application/pdf") == "pdf"
    assert (
        detect_format(
            dummy_bytes,
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )
        == "docx"
    )


def test_detect_format_by_filename() -> None:
    dummy_bytes = b"non-magic-content"
    assert detect_format(dummy_bytes, filename="course_syllabus.pdf") == "pdf"
    assert detect_format(dummy_bytes, filename="course_syllabus.docx") == "docx"


def test_detect_format_unsupported() -> None:
    with pytest.raises(UnsupportedFormatError, match="Unsupported or unrecognized"):
        detect_format(b"Plain text file content", filename="notes.txt")


def test_parse_document_pdf_end_to_end() -> None:
    pdf_bytes = _make_pdf_bytes("Intro to Machine Learning")
    extracted = parse_document(pdf_bytes, filename="ml.pdf")
    assert "Intro to Machine Learning" in extracted


def test_parse_document_docx_end_to_end() -> None:
    docx_bytes = _make_docx_bytes("Advanced Algorithms & Data Structures")
    extracted = parse_document(docx_bytes, filename="algo.docx")
    assert "Advanced Algorithms & Data Structures" in extracted


def test_parse_document_with_cleaning() -> None:
    doc = docx.Document()
    doc.add_paragraph("• Algo-\nrithms and Complexity")
    doc.add_paragraph("Page 1 of 3")
    buf = io.BytesIO()
    doc.save(buf)

    cleaned = parse_document(buf.getvalue(), filename="cs.docx", clean=True)
    assert "- Algorithms and Complexity" in cleaned
    assert "Page 1 of 3" not in cleaned


def test_parse_document_unsupported_format() -> None:
    with pytest.raises(UnsupportedFormatError):
        parse_document(b"Some random binary", filename="archive.tar.gz")
