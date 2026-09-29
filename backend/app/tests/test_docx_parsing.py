import io

import docx
import pytest

from app.services.parsing.docx import (
    DOCXParsingError,
    DOCXSizeLimitExceededError,
    EmptyDOCXError,
    extract_text_from_docx,
)


def _generate_docx_bytes(
    paragraphs: list[str] | None = None,
    table_rows: list[list[str]] | None = None,
) -> bytes:
    """Helper to generate a valid in-memory DOCX document."""
    doc = docx.Document()
    if paragraphs:
        for text in paragraphs:
            doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r_idx, row in enumerate(table_rows):
            for c_idx, val in enumerate(row):
                table.cell(r_idx, c_idx).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_docx_paragraphs() -> None:
    data = _generate_docx_bytes(
        paragraphs=[
            "CS 301: Cloud Infrastructure",
            "Topics include Docker, K8s, and Terraform.",
        ]
    )
    text = extract_text_from_docx(data)
    assert "CS 301: Cloud Infrastructure" in text
    assert "Topics include Docker, K8s, and Terraform." in text


def test_extract_docx_tables() -> None:
    table_data = [
        ["Week", "Topic", "Lab"],
        ["Week 1", "Container Basics", "Docker run"],
        ["Week 2", "Orchestration", "Kubernetes Pods"],
    ]
    data = _generate_docx_bytes(
        paragraphs=["Course Schedule:"],
        table_rows=table_data,
    )
    text = extract_text_from_docx(data)
    assert "Course Schedule:" in text
    assert "Container Basics" in text
    assert "Kubernetes Pods" in text
    assert "Docker run" in text


def test_extract_docx_from_stream() -> None:
    data = _generate_docx_bytes(paragraphs=["Stream extraction test."])
    stream = io.BytesIO(data)
    text = extract_text_from_docx(stream)
    assert "Stream extraction test." in text


def test_extract_docx_empty_source() -> None:
    with pytest.raises(EmptyDOCXError, match="DOCX source is empty"):
        extract_text_from_docx(b"")


def test_extract_docx_no_text() -> None:
    doc = docx.Document()
    buf = io.BytesIO()
    doc.save(buf)
    with pytest.raises(EmptyDOCXError, match="does not contain any extractable text"):
        extract_text_from_docx(buf.getvalue())


def test_extract_docx_size_limit() -> None:
    data = _generate_docx_bytes(paragraphs=["Sample content"])
    with pytest.raises(DOCXSizeLimitExceededError, match="exceeds limit"):
        extract_text_from_docx(data, max_size=10)


def test_extract_docx_corrupted_data() -> None:
    with pytest.raises(DOCXParsingError, match="Failed to open DOCX"):
        extract_text_from_docx(b"PK\x03\x04corrupted-not-a-real-docx")
